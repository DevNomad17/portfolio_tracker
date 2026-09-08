"""Read deposits from rodina_portfolia.xlsm (sheet per portfolio: Trade Date YYYYMMDD, date, note, EUR)."""
from dataclasses import dataclass
from datetime import date, datetime
import openpyxl

@dataclass(frozen=True)
class Deposit:
    date: date
    note: str
    amount: float

def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(int(v)) if isinstance(v, (int, float)) else str(v).strip()
    return datetime.strptime(s, "%Y%m%d").date()

def read_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not in {path} (have {wb.sheetnames})")
    ws = wb[sheet]
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c in (None, "") for c in row[:4]):
            continue
        trade, dt, note, amt = (list(row) + [None]*4)[:4]
        if amt in (None, ""):
            raise ValueError(f"{sheet} row {i}: missing amount")
        amt = float(amt)
        if amt <= 0:
            raise ValueError(f"{sheet} row {i}: amount must be positive, got {amt}")
        d = _to_date(dt if dt not in (None, "") else trade)
        if trade not in (None, "") and _to_date(trade) != d:
            raise ValueError(f"{sheet} row {i}: Trade Date {trade} != date {dt}")
        if d > date.today():
            raise ValueError(f"{sheet} row {i}: deposit date {d} is in the future")
        out.append(Deposit(d, (note or "").strip(), amt))
    wb.close()
    if not out:
        raise ValueError(f"{sheet}: no deposits found")
    return sorted(out, key=lambda x: x.date)

def read_all(cfg):
    """{person_id: [Deposit,...]} — people sharing a sheet each get the full list."""
    from . import abspath
    path = abspath(cfg, cfg["workbook"])
    cache = {}
    res = {}
    for pid, p in cfg["people"].items():
        sh = p["sheet"]
        if sh not in cache:
            cache[sh] = read_sheet(path, sh)
        res[pid] = list(cache[sh])
    return res
